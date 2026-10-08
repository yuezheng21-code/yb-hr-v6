"""
渊博579 HR V7 — 文书模板渲染

模板语法（纯文本，易于 HR 编辑）：
  # 标题 / ## 小节标题
  空行分段；**粗体**；--- 分页
  {{key}}            自动填充（见 AUTO_FIELDS），未知的 key 成为需要 HR 填写的字段
  {{key|标签}}       同上，并给出填写表单里显示的标签
输出：HTML 预览、DOCX（可在 Word 中继续修改）。
"""
from __future__ import annotations
import html
import io
import re
from datetime import date
from typing import Optional

PLACEHOLDER = re.compile(r"\{\{\s*([a-zA-Z0-9_.]+)\s*(?:\|\s*([^}]*?)\s*)?\}\}")

AUTO_FIELDS = {
    "employee.name": "员工姓名", "employee.emp_no": "工号", "employee.birth_date": "出生日期",
    "employee.nationality": "国籍", "employee.address": "住址", "employee.phone": "电话",
    "employee.tax_id": "税号", "employee.social_security_no": "社保号", "employee.iban": "IBAN",
    "employee.position": "岗位", "employee.grade": "职级", "employee.hourly_rate": "当前时薪",
    "employee.primary_warehouse": "主仓库", "employee.join_date": "入职日期", "employee.leave_date": "离职日期",
    "contract.contract_no": "合同编号", "contract.start_date": "合同开始", "contract.end_date": "合同结束",
    "contract.probation_end": "试用期结束", "contract.weekly_hours": "每周工时", "contract.hourly_rate": "合同时薪",
    "contract.vacation_days": "年假天数", "contract.position": "合同岗位", "contract.warehouse_code": "工作地点（仓库）",
    "contract.notice_period": "解约期", "contract.signed_date": "签署日期",
    "company.name": "公司名称", "company.address": "公司地址", "company.representative": "公司代表",
    "today": "今天日期",
}


def fmt_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, float):
        s = f"{v:,.2f}" if v % 1 else f"{int(v):,}"
        return s.replace(",", "X").replace(".", ",").replace("X", ".")  # German number format
    return str(v)


def fields_of(body: str) -> list[dict]:
    """Ordered unique placeholders: [{key, label, auto}]"""
    seen, out = set(), []
    for m in PLACEHOLDER.finditer(body or ""):
        key, label = m.group(1), m.group(2)
        if key in seen:
            continue
        seen.add(key)
        out.append({"key": key, "label": label or AUTO_FIELDS.get(key, key), "auto": key in AUTO_FIELDS,
                    "optional": _optional(label)})
    return out


def build_context(employee, contract, settings: dict) -> dict:
    ctx = {"today": fmt_value(date.today())}
    for k in ("name", "emp_no", "birth_date", "nationality", "address", "phone", "tax_id", "social_security_no",
              "iban", "position", "grade", "hourly_rate", "primary_warehouse", "join_date", "leave_date"):
        ctx[f"employee.{k}"] = fmt_value(getattr(employee, k, None)) if employee is not None else ""
    for k in ("contract_no", "start_date", "end_date", "probation_end", "weekly_hours", "hourly_rate",
              "vacation_days", "position", "warehouse_code", "notice_period", "signed_date"):
        ctx[f"contract.{k}"] = fmt_value(getattr(contract, k, None)) if contract is not None else ""
    ctx["company.name"] = settings.get("company_name", "")
    ctx["company.address"] = settings.get("company_address", "")
    ctx["company.representative"] = settings.get("company_representative", "")
    return ctx


def _optional(label) -> bool:
    """{{key|… (optional)}} / （可选）: may stay empty."""
    return bool(label) and ("optional" in label.lower() or "可选" in label)


def fill(body: str, ctx: dict, values: dict) -> tuple[str, list[str]]:
    """Replace placeholders. Explicit values override auto values. Returns (text, missing_keys)."""
    missing: list[str] = []

    def rep(m):
        key = m.group(1)
        v = values.get(key)
        if v in (None, ""):
            v = ctx.get(key, "")
        if v in (None, ""):
            if _optional(m.group(2)):
                return ""
            missing.append(key)
            return f"[{m.group(2) or AUTO_FIELDS.get(key, key)}]"
        return str(v)

    return PLACEHOLDER.sub(rep, body or ""), sorted(set(missing))


def _blocks(text: str):
    """Yield (kind, content): h1/h2/p/pagebreak."""
    para: list[str] = []
    for raw in text.splitlines():
        line = raw.rstrip()
        if line.strip() == "---":
            if para:
                yield "p", "\n".join(para); para = []
            yield "pagebreak", ""
        elif line.startswith("## "):
            if para:
                yield "p", "\n".join(para); para = []
            yield "h2", line[3:].strip()
        elif line.startswith("# "):
            if para:
                yield "p", "\n".join(para); para = []
            yield "h1", line[2:].strip()
        elif not line.strip():
            if para:
                yield "p", "\n".join(para); para = []
        else:
            para.append(line)
    if para:
        yield "p", "\n".join(para)


_BOLD = re.compile(r"\*\*(.+?)\*\*")


def to_html(text: str) -> str:
    out = []
    for kind, content in _blocks(text):
        esc = _BOLD.sub(r"<b>\1</b>", html.escape(content)).replace("\n", "<br>")
        if kind == "h1":
            out.append(f"<h1>{esc}</h1>")
        elif kind == "h2":
            out.append(f"<h2>{esc}</h2>")
        elif kind == "pagebreak":
            out.append('<hr class="pb">')
        else:
            out.append(f"<p>{esc}</p>")
    return "\n".join(out)


def to_docx(text: str, title: Optional[str] = None) -> bytes:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt, Cm

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)
    if title:
        doc.core_properties.title = title
    for kind, content in _blocks(text):
        if kind == "h1":
            doc.add_heading(content, level=1)
        elif kind == "h2":
            doc.add_heading(content, level=2)
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            p = doc.add_paragraph()
            lines = content.split("\n")
            for i, line in enumerate(lines):
                pos = 0
                for m in _BOLD.finditer(line):
                    if m.start() > pos:
                        p.add_run(line[pos:m.start()])
                    p.add_run(m.group(1)).bold = True
                    pos = m.end()
                if pos < len(line):
                    p.add_run(line[pos:])
                if i < len(lines) - 1:
                    p.add_run().add_break()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
