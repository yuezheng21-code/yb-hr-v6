"""
渊博579 HR V7 — 作业数据接入服务

把任意来源（手工 / 扫码 / Excel / WMS API 推送）的作业记录规范化为 OperationLog：
  1. 字段规范化：各系统导出列名 → 标准字段（中/英/德列名同义词 + 系统预设 + 自定义映射）
  2. 员工匹配：工号 → 操作员映射(OperatorAlias) → 工号=账号 → 姓名
  3. 作业类型匹配：代码 → 名称 → 关键词
  4. 幂等：同一来源的 external_ref 只入库一次，WMS 重复推送不会重复计数
未识别员工/作业类型的记录以 status=unmatched 保存，补充映射后可一键重新匹配。

注意：马帮/领星/易仓的预设列名基于常见导出字段整理，各账号的导出模板可自定义，
      首次接入时请用 dry_run 预览核对，必要时在接入源上配置 field_mapping 覆盖。
"""
from __future__ import annotations
import csv
import io
import json
from datetime import date, datetime, timedelta
from typing import Any, Iterable, Optional
from sqlalchemy import select, func
from sqlalchemy.orm import Session
from backend.models.employee import Employee
from backend.models.operation import OperationLog, OperationType, OperatorAlias
from backend.services.sequence import next_sequence_no, make_prefix
from backend.services.settlement_calc import compute_hours

# ── 标准字段 → 常见列名同义词 ────────────────────────────────────────────
FIELD_ALIASES: dict[str, list[str]] = {
    "external_ref": ["external_ref", "id", "task_id", "record_id", "log_id", "作业单号", "任务号", "任务ID",
                     "流水号", "记录ID", "单据编号"],
    "emp_no": ["emp_no", "employee_no", "工号", "员工编号", "员工工号", "Personalnummer"],
    "operator": ["operator", "operator_account", "operator_name", "user", "username", "account", "操作人",
                 "操作员", "操作人员", "作业人", "作业人员", "拣货员", "拣货人", "打包员", "打包人", "上架人",
                 "收货人", "员工", "员工姓名", "姓名", "Mitarbeiter", "Bearbeiter"],
    "work_date": ["work_date", "date", "日期", "作业日期", "操作日期", "业务日期", "Datum"],
    "timestamp": ["timestamp", "time", "datetime", "操作时间", "作业时间", "完成时间", "扫描时间", "创建时间",
                  "Zeitpunkt"],
    "op": ["op_code", "op", "operation", "op_type", "task_type", "作业类型", "操作类型", "任务类型", "作业环节",
           "作业节点", "操作节点", "工序", "Tätigkeit", "Vorgang"],
    "qty": ["qty", "quantity", "pieces", "count", "数量", "件数", "商品数量", "作业数量", "产品数量", "SKU数量",
            "Menge", "Stück"],
    "error_qty": ["error_qty", "errors", "差错数", "异常数量", "错误数量", "Fehler"],
    "hours": ["hours", "工时", "时长", "作业时长(小时)", "Stunden"],
    "start_time": ["start_time", "开始时间", "Beginn"],
    "end_time": ["end_time", "结束时间", "Ende"],
    "warehouse_code": ["warehouse_code", "warehouse", "仓库", "仓库代码", "仓库编码", "Lager"],
    "client": ["client", "customer", "客户", "客户名称", "客户代码", "货主", "Kunde"],
    "ref_no": ["ref_no", "order_no", "单号", "订单号", "出库单号", "入库单号", "波次号", "拣货单号", "柜号",
               "container_no", "项目号", "Auftrag"],
    "notes": ["notes", "remark", "备注", "Bemerkung"],
}

# ── 系统预设（追加的列名同义词） ─────────────────────────────────────────
SYSTEM_PRESETS: dict[str, dict[str, Any]] = {
    "generic": {
        "label": "通用模板 / 本系统模板",
        "aliases": {},
    },
    "mabang": {
        "label": "马帮 ERP/WMS",
        "aliases": {
            "operator": ["配货员", "复核员", "包装员", "称重员"],
            "op": ["作业类型名称"],
            "ref_no": ["包裹号", "马帮订单号"],
            "qty": ["配货数量", "复核数量"],
        },
    },
    "lingxing": {
        "label": "领星 ERP / 海外仓 WMS",
        "aliases": {
            "operator": ["操作账号", "作业人账号"],
            "ref_no": ["出库单号", "入库单号", "货件单号", "Shipment ID", "FBA号"],
            "client": ["客户简称"],
            "qty": ["实际数量", "完成数量"],
        },
    },
    "eccang": {
        "label": "易仓 WMS/OMS",
        "aliases": {
            "operator": ["操作人账号", "作业员"],
            "ref_no": ["订单号(参考号)", "参考号", "跟踪号"],
            "client": ["客户代码"],
            "qty": ["操作数量"],
        },
    },
}

DEFAULT_SYSTEMS = list(SYSTEM_PRESETS.keys())


def _norm_key(k: Any) -> str:
    return str(k or "").strip().lower().replace(" ", "").replace("（", "(").replace("）", ")")


def build_header_map(headers: Iterable[str], system: str = "generic",
                     override: Optional[dict] = None) -> dict[str, str]:
    """Return {canonical_field: source_header} for the given source headers."""
    headers = [h for h in headers if h is not None]
    by_norm = {_norm_key(h): h for h in headers}
    result: dict[str, str] = {}
    if override:
        for field, src in override.items():
            if src and _norm_key(src) in by_norm:
                result[field] = by_norm[_norm_key(src)]
    extra = SYSTEM_PRESETS.get(system, {}).get("aliases", {})
    for field, aliases in FIELD_ALIASES.items():
        if field in result:
            continue
        for a in list(extra.get(field, [])) + aliases:
            if _norm_key(a) in by_norm:
                result[field] = by_norm[_norm_key(a)]
                break
    return result


# ── 值解析 ───────────────────────────────────────────────────────────────
def parse_number(v: Any, default: float = 0.0) -> float:
    if v is None or v == "":
        return default
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace("€", "")
    if "," in s and "." in s:  # 1.234,5 (DE) or 1,234.5 (EN)
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return default


def parse_datetime(v: Any) -> Optional[datetime]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    if isinstance(v, (int, float)) and 20000 < float(v) < 80000:  # Excel serial date
        return datetime(1899, 12, 30) + timedelta(days=float(v))
    s = str(v).strip().replace("T", " ").split("+")[0].rstrip("Z")
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M",
                "%Y/%m/%d", "%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S.%f"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def parse_hhmm(v: Any) -> Optional[str]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.strftime("%H:%M")
    if hasattr(v, "strftime"):
        return v.strftime("%H:%M")
    s = str(v).strip()
    dt = parse_datetime(s)
    if dt and (" " in s or "T" in s):
        return dt.strftime("%H:%M")
    parts = s.split(":")
    if len(parts) >= 2 and parts[0].isdigit() and parts[1][:2].isdigit():
        return f"{int(parts[0]):02d}:{int(parts[1][:2]):02d}"
    return None


def normalize_record(raw: dict, header_map: dict[str, str]) -> dict:
    """Map a raw source row to canonical fields."""
    def g(field):
        src = header_map.get(field)
        return raw.get(src) if src is not None else None

    ts = parse_datetime(g("timestamp"))
    wd = parse_datetime(g("work_date")) or ts
    rec = {
        "external_ref": str(g("external_ref")).strip() if g("external_ref") not in (None, "") else None,
        "emp_no": str(g("emp_no")).strip() if g("emp_no") not in (None, "") else None,
        "operator": str(g("operator")).strip() if g("operator") not in (None, "") else None,
        "work_date": wd.date() if wd else None,
        "op": str(g("op")).strip() if g("op") not in (None, "") else None,
        "qty": parse_number(g("qty"), default=1.0),
        "error_qty": parse_number(g("error_qty")),
        "hours": parse_number(g("hours")),
        "start_time": parse_hhmm(g("start_time")),
        "end_time": parse_hhmm(g("end_time")),
        "warehouse_code": str(g("warehouse_code")).strip() if g("warehouse_code") not in (None, "") else None,
        "client": str(g("client")).strip() if g("client") not in (None, "") else None,
        "ref_no": str(g("ref_no")).strip()[:60] if g("ref_no") not in (None, "") else None,
        "notes": str(g("notes")).strip() if g("notes") not in (None, "") else None,
    }
    if ts and not rec["start_time"] and not rec["end_time"] and " " in str(g("timestamp") or ""):
        rec["end_time"] = ts.strftime("%H:%M")
    return rec


def read_tabular_file(filename: str, content: bytes) -> list[dict]:
    """Read CSV or XLSX into a list of row dicts (first row = header)."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = None
        out = []
        for row in it:
            if header is None:
                if row and any(c not in (None, "") for c in row):
                    header = [str(c).strip() if c is not None else "" for c in row]
                continue
            if not row or all(c in (None, "") for c in row):
                continue
            out.append({header[i]: row[i] for i in range(min(len(header), len(row))) if header[i]})
        wb.close()
        return out
    for enc in ("utf-8-sig", "gb18030", "latin-1"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:4096]
    delim = ";" if sample.count(";") > sample.count(",") else ("\t" if sample.count("\t") > sample.count(",") else ",")
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)
    return [{(k or "").strip(): v for k, v in row.items()} for row in reader if any((v or "").strip() for v in row.values())]


# ── 匹配 ─────────────────────────────────────────────────────────────────
class Matcher:
    """Per-batch cache for employee / operation-type resolution."""

    def __init__(self, db: Session, system: str = "generic"):
        self.db = db
        self.system = system
        self._emp_cache: dict[tuple, Optional[Employee]] = {}
        self.op_types = db.scalars(select(OperationType).where(OperationType.is_active == True)).all()  # noqa: E712

    def employee(self, emp_no: Optional[str], operator: Optional[str]) -> Optional[Employee]:
        key = (emp_no, operator)
        if key in self._emp_cache:
            return self._emp_cache[key]
        db, emp = self.db, None
        if emp_no:
            emp = db.scalar(select(Employee).where(Employee.emp_no == emp_no))
        if emp is None and operator:
            alias = db.scalar(select(OperatorAlias).where(
                OperatorAlias.alias == operator, OperatorAlias.system.in_([self.system, "*"])
            ).order_by(OperatorAlias.system.desc()))
            if alias:
                emp = db.get(Employee, alias.employee_id)
            if emp is None:
                emp = db.scalar(select(Employee).where(Employee.emp_no == operator))
            if emp is None:
                emp = db.scalar(select(Employee).where(func.lower(Employee.name) == operator.lower()))
        self._emp_cache[key] = emp
        return emp

    def op_type(self, op: Optional[str], client: Optional[str] = None) -> Optional[OperationType]:
        if not op:
            return None
        s = op.strip().lower()
        candidates = sorted(self.op_types, key=lambda t: 0 if (client and t.client and t.client.lower() == client.lower()) else 1)
        for t in candidates:
            if t.code.lower() == s:
                return t
        for t in candidates:
            if t.name.lower() == s or (t.name_de and t.name_de.lower() == s):
                return t
        for t in candidates:
            for kw in (t.keywords or "").split(","):
                kw = kw.strip().lower()
                if kw and kw in s:
                    return t
        return None


def next_op_no(db: Session) -> str:
    return next_sequence_no(db, OperationLog, OperationLog.op_no, make_prefix("OP"), width=6)


def apply_match(log: OperationLog, matcher: Matcher) -> None:
    """(Re)resolve employee and operation type; set status/match_issue accordingly."""
    issues = []
    if log.employee_id is None:
        emp = matcher.employee(log.emp_no, log.operator_ref)
        if emp:
            log.employee_id, log.emp_no, log.emp_name, log.supplier_id = emp.id, emp.emp_no, emp.name, emp.supplier_id
        else:
            issues.append(f"员工未识别: {log.operator_ref or log.emp_no or '-'}")
    if log.op_type_id is None:
        ot = matcher.op_type(log.op_code or log.op_label, log.client)
        if ot:
            log.op_type_id, log.op_code = ot.id, ot.code
            log.client = log.client or ot.client
        else:
            issues.append(f"作业类型未识别: {log.op_label or log.op_code or '-'}")
    if issues:
        log.status = "unmatched"
        log.match_issue = "; ".join(issues)[:200]
    elif log.status == "unmatched":
        log.status = "pending"
        log.match_issue = None


def ingest_records(
    db: Session,
    records: list[dict],
    *,
    source: str,
    system: str = "generic",
    field_mapping: Optional[dict] = None,
    default_warehouse: Optional[str] = None,
    default_client: Optional[str] = None,
    default_date: Optional[date] = None,
    auto_confirm: bool = False,
    created_by: str = "",
    dry_run: bool = False,
) -> dict:
    """Normalize + match + persist a batch. Caller commits."""
    headers: list[str] = []
    for r in records[:50]:
        for k in r.keys():
            if k not in headers:
                headers.append(k)
    header_map = build_header_map(headers, system, field_mapping)
    matcher = Matcher(db, system)

    created, duplicates, unmatched, errors, preview = 0, 0, 0, [], []
    seen_refs: set = set()
    next_no: Optional[str] = None
    for idx, raw in enumerate(records, start=1):
        rec = normalize_record(raw, header_map)
        wd = rec["work_date"] or default_date
        if wd is None:
            errors.append({"row": idx, "error": "缺少日期"})
            continue
        if not (rec["emp_no"] or rec["operator"]):
            errors.append({"row": idx, "error": "缺少操作人/工号"})
            continue
        ext = f"{system}:{rec['external_ref']}" if rec["external_ref"] else None
        if ext:
            if ext in seen_refs or db.scalar(select(OperationLog.id).where(
                    OperationLog.source == source, OperationLog.external_ref == ext)):
                duplicates += 1
                continue
            seen_refs.add(ext)

        hours = rec["hours"]
        if not hours and rec["start_time"] and rec["end_time"]:
            hours = compute_hours(rec["start_time"], rec["end_time"])
        log = OperationLog(
            op_no="",
            emp_no=rec["emp_no"],
            operator_ref=rec["operator"] or rec["emp_no"],
            op_label=rec["op"],
            op_code=rec["op"][:30] if rec["op"] else None,
            work_date=wd,
            warehouse_code=(rec["warehouse_code"] or default_warehouse or None),
            client=rec["client"] or default_client,
            ref_no=rec["ref_no"],
            qty=rec["qty"],
            error_qty=rec["error_qty"],
            start_time=rec["start_time"],
            end_time=rec["end_time"],
            hours=hours or 0.0,
            source=source,
            source_system=system,
            external_ref=ext,
            status="pending",
            created_by=created_by,
            notes=rec["notes"],
        )
        apply_match(log, matcher)
        if log.status == "unmatched":
            unmatched += 1
        elif auto_confirm:
            log.status = "confirmed"
            log.confirmed_by = created_by or "auto"
            log.confirmed_at = datetime.utcnow()
        if dry_run:
            if len(preview) < 20:
                preview.append({
                    "row": idx, "work_date": str(wd), "operator": log.operator_ref, "emp_name": log.emp_name,
                    "op_label": log.op_label, "op_code": log.op_code, "qty": log.qty, "hours": log.hours,
                    "ref_no": log.ref_no, "status": log.status, "match_issue": log.match_issue,
                })
        else:
            if next_no is None:
                next_no = next_op_no(db)
            log.op_no = next_no
            prefix, seq = next_no.rsplit("-", 1)
            next_no = f"{prefix}-{int(seq) + 1:0{len(seq)}d}"
            db.add(log)
        created += 1

    return {
        "total": len(records),
        "created": 0 if dry_run else created,
        "would_create": created if dry_run else None,
        "duplicates": duplicates,
        "unmatched": unmatched,
        "errors": errors[:100],
        "header_map": header_map,
        "preview": preview if dry_run else None,
    }


def parse_mapping(s: Optional[str]) -> Optional[dict]:
    if not s:
        return None
    try:
        v = json.loads(s)
        return v if isinstance(v, dict) else None
    except json.JSONDecodeError:
        return None
