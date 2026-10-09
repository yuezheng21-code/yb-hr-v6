"""
渊博579 HR V7 — 通用列表导出 /api/v1/export

POST /table?format=xlsx|pdf
  前端把「当前筛选后的列表」按列定义发来，服务端渲染成 Excel 或 PDF（横向 A4）。
  只渲染调用者已经能看到的数据，不额外查询数据库；甲方账号不可用。
"""
from __future__ import annotations
import io
from datetime import date, datetime
from typing import Any, Optional
from urllib.parse import quote
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle
from backend.middleware.auth import get_current_user
from backend.models.user import User
from backend.services import settings_store
from backend.services.finance_pdf import P, eur, num, de_date, STY, LINE, SOFT

router = APIRouter(prefix="/api/v1/export", tags=["export"])

MAX_ROWS, MAX_COLS = 20000, 40
TYPES = {"text", "num", "int", "money", "date", "pct"}


class Column(BaseModel):
    label: str = Field(..., max_length=80)
    type: str = "text"
    sum: bool = False


class TableIn(BaseModel):
    title: str = Field(..., max_length=120)
    subtitle: Optional[str] = Field(None, max_length=500)
    columns: list[Column]
    rows: list[list[Any]]


def _num(v) -> Optional[float]:
    if v in (None, ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _date(v) -> Optional[date]:
    if not v:
        return None
    try:
        return date.fromisoformat(str(v)[:10])
    except ValueError:
        return None


def _fname(title: str, ext: str) -> str:
    return f"{title}_{datetime.now():%Y%m%d_%H%M}.{ext}"


def _xlsx(body: TableIn, user: User) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = body.title[:31].replace("/", "-") or "Export"
    ncol = len(body.columns)
    ws.cell(1, 1, body.title).font = Font(bold=True, size=14)
    ws.cell(2, 1, " · ".join(x for x in [body.subtitle, f"导出：{user.display_name} {datetime.now():%Y-%m-%d %H:%M}"] if x)).font = Font(color="666666", size=9)
    hdr = 4
    fill = PatternFill("solid", fgColor="F1F3F5")
    for j, c in enumerate(body.columns, 1):
        cell = ws.cell(hdr, j, c.label)
        cell.font, cell.fill = Font(bold=True), fill
        cell.alignment = Alignment(horizontal="right" if c.type in ("num", "int", "money", "pct") else "left", vertical="center")
    widths = [max(8, min(len(c.label) * 2, 40)) for c in body.columns]
    for i, row in enumerate(body.rows, hdr + 1):
        for j, c in enumerate(body.columns, 1):
            v = row[j - 1] if j - 1 < len(row) else None
            cell = ws.cell(i, j)
            if c.type in ("num", "int", "money", "pct"):
                n = _num(v)
                cell.value = n if n is not None else (str(v) if v not in (None, "") else None)
                cell.number_format = {"money": '#,##0.00 "€"', "int": "#,##0", "pct": '0.0"%"'}.get(c.type, "#,##0.00")
            elif c.type == "date":
                d = _date(v)
                cell.value = d if d else (str(v) if v else None)
                cell.number_format = "DD.MM.YYYY"
            else:
                cell.value = "" if v is None else str(v)[:1000]
            widths[j - 1] = max(widths[j - 1], min(len(str(v or "")) + 2, 60))
    last = hdr + len(body.rows)
    if body.rows and any(c.sum for c in body.columns):
        t = last + 1
        ws.cell(t, 1, "合计").font = Font(bold=True)
        for j, c in enumerate(body.columns, 1):
            if c.sum:
                col = get_column_letter(j)
                cell = ws.cell(t, j, f"=SUBTOTAL(9,{col}{hdr + 1}:{col}{last})")
                cell.font = Font(bold=True)
                cell.number_format = {"money": '#,##0.00 "€"', "int": "#,##0"}.get(c.type, "#,##0.00")
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(hdr + 1, 1)
    if body.rows:
        ws.auto_filter.ref = f"A{hdr}:{get_column_letter(ncol)}{last}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _fmt(v, t: str) -> str:
    if v in (None, ""):
        return ""
    if t == "money":
        n = _num(v)
        return eur(n) if n is not None else str(v)
    if t in ("num", "pct"):
        n = _num(v)
        return (num(n) + ("%" if t == "pct" else "")) if n is not None else str(v)
    if t == "int":
        n = _num(v)
        return f"{int(round(n)):,}".replace(",", ".") if n is not None else str(v)
    if t == "date":
        d = _date(v)
        return de_date(d) if d else str(v)
    return str(v)[:300]


def _pdf(body: TableIn, user: User) -> bytes:
    company = settings_store.get("company_name") or ""
    cols = body.columns
    n = len(cols)
    size = 8 if n <= 8 else 7 if n <= 12 else 6
    from reportlab.lib.styles import ParagraphStyle
    cell = ParagraphStyle("c", parent=STY["base"], fontSize=size, leading=size + 2)
    cellr = ParagraphStyle("cr", parent=cell, alignment=2)
    head = ParagraphStyle("h", parent=cell, fontName="Helvetica-Bold")
    headr = ParagraphStyle("hr", parent=head, alignment=2)
    STY.update({"_c": cell, "_cr": cellr, "_h": head, "_hr": headr})
    right = [c.type in ("num", "int", "money", "pct") for c in cols]
    data = [[P(c.label, "_hr" if right[j] else "_h") for j, c in enumerate(cols)]]
    lens = [len(c.label) for c in cols]
    for row in body.rows:
        line = []
        for j, c in enumerate(cols):
            s = _fmt(row[j] if j < len(row) else None, c.type)
            lens[j] = max(lens[j], min(len(s), 40))
            line.append(P(s, "_cr" if right[j] else "_c"))
        data.append(line)
    if body.rows and any(c.sum for c in cols):
        tot = []
        for j, c in enumerate(cols):
            if c.sum:
                s = sum(_num(r[j]) or 0 for r in body.rows if j < len(r))
                tot.append(P(f"**{_fmt(s, c.type)}**", "_cr"))
            else:
                tot.append(P("**合计**" if j == 0 else "", "_c"))
        data.append(tot)
    avail = landscape(A4)[0] - 24 * mm
    weights = [max(4, l) for l in lens]
    widths = [avail * w / sum(weights) for w in weights]
    t = Table(data, colWidths=widths, repeatRows=1)
    st = [("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (-1, 0), SOFT),
          ("LINEBELOW", (0, 0), (-1, -1), 0.25, LINE), ("TOPPADDING", (0, 0), (-1, -1), 2),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]
    if body.rows and any(c.sum for c in cols):
        st.append(("LINEABOVE", (0, -1), (-1, -1), 0.8, colors.black))
    t.setStyle(TableStyle(st))

    def deco(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#6b7280"))
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 8 * mm, f"{doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm,
                            bottomMargin=14 * mm, title=body.title)
    sub = " · ".join(x for x in [company, body.subtitle, f"{len(body.rows)} 条", f"{user.display_name} {datetime.now():%d.%m.%Y %H:%M}"] if x)
    doc.build([P(body.title, "title"), P(sub, "small"), Spacer(1, 4 * mm), t], onFirstPage=deco, onLaterPages=deco)
    return buf.getvalue()


@router.post("/table")
def export_table(body: TableIn, format: str = Query("xlsx", pattern="^(xlsx|pdf)$"), user: User = Depends(get_current_user)):
    if user.role == "client":
        raise HTTPException(403, "Forbidden")
    if not body.columns or len(body.columns) > MAX_COLS:
        raise HTTPException(400, f"列数应为 1–{MAX_COLS}")
    if len(body.rows) > MAX_ROWS:
        raise HTTPException(413, f"单次最多导出 {MAX_ROWS} 行，请缩小筛选范围")
    for c in body.columns:
        if c.type not in TYPES:
            c.type = "text"
    if format == "xlsx":
        data, media, ext = _xlsx(body, user), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    else:
        data, media, ext = _pdf(body, user), "application/pdf", "pdf"
    return Response(data, media_type=media, headers={
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(_fname(body.title, ext))}",
        "Cache-Control": "private, no-store"})
